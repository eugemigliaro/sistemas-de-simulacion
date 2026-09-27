#!/usr/bin/env python3
"""Arma el PPTX de la exposición oral a partir del PDF, con los videos embebidos.

La guía de presentaciones exige que, en vivo, las animaciones estén embebidas
en la diapositiva, mientras que el PDF entregado lleva un fotograma y el link
[GPRES, p. 3]. Cada diapositiva del PPTX es la página del PDF como imagen, y
sobre las que tienen animación se superpone el MP4 exactamente donde está el
fotograma.

Para ubicar los videos no se estiman medidas de Beamer: se compila una copia
de la presentación con cada fotograma reemplazado por un rectángulo de color
propio y se busca ese color en las páginas renderizadas.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
PRESENTACION = ROOT / "presentacion"
STEM = "SdS_TP3_2026Q2G07CS_Presentacion"
TECTONIC = shutil.which("tectonic") or str(
    Path.home() / ".codex/.tmp/bundled-marketplaces/openai-bundled/plugins/latex/bin/tectonic")
ANCHO, ALTO = 1920, 1080

# Fotograma del PDF -> video y color de marcado.
VIDEOS = {
    "anim-goles-vacia.png": (ROOT / "videos/anim-goles-vacia.mp4", (255, 0, 255)),
    "anim-goles-elegida.png": (ROOT / "videos/anim-goles-elegida.mp4", (0, 255, 255)),
}


def renderizar(pdf: Path, destino: Path) -> list[Path]:
    destino.mkdir(parents=True, exist_ok=True)
    subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(ANCHO), "-scale-to-y", str(ALTO),
                    str(pdf), str(destino / "p")], check=True)
    return sorted(destino.glob("p-*.png"))


def ubicar_fotogramas(trabajo: Path) -> dict[int, tuple[str, tuple[float, float, float, float]]]:
    """Página (desde 0) -> (fotograma, caja en fracciones de la página)."""
    copia = trabajo / "mascara"
    shutil.copytree(PRESENTACION, copia, ignore=shutil.ignore_patterns("*.pptx", f"{STEM}.pdf"))
    for nombre, (_, color) in VIDEOS.items():
        with Image.open(PRESENTACION / "figuras" / nombre) as original:
            Image.new("RGB", original.size, color).save(copia / "figuras" / nombre)
    subprocess.run([TECTONIC, "--only-cached", f"{STEM}.tex"], cwd=copia, check=True,
                   capture_output=True, text=True)
    cajas = {}
    for indice, pagina in enumerate(renderizar(copia / f"{STEM}.pdf", trabajo / "mascara-png")):
        pixeles = np.asarray(Image.open(pagina).convert("RGB")).astype(int)
        for nombre, (_, color) in VIDEOS.items():
            coincide = np.all(np.abs(pixeles - np.array(color)) < 12, axis=2)
            if coincide.sum() < 1000:
                continue
            filas, columnas = np.nonzero(coincide)
            cajas[indice] = (nombre, (columnas.min() / ANCHO, filas.min() / ALTO,
                                      (columnas.max() + 1) / ANCHO, (filas.max() + 1) / ALTO))
    return cajas


def main() -> int:
    pdf = PRESENTACION / f"{STEM}.pdf"
    salida = PRESENTACION / f"{STEM}.pptx"
    for video, _ in VIDEOS.values():
        if not video.exists():
            raise FileNotFoundError(f"falta {video}: correr barridos_presentacion.py animaciones")
    with tempfile.TemporaryDirectory() as directorio:
        trabajo = Path(directorio)
        paginas = renderizar(pdf, trabajo / "paginas")
        cajas = ubicar_fotogramas(trabajo)
        faltan = set(VIDEOS) - {nombre for nombre, _ in cajas.values()}
        if faltan:
            raise RuntimeError(f"no se encontró en el PDF el fotograma de {sorted(faltan)}")

        presentacion = Presentation()
        presentacion.slide_width = Emu(12192000)   # 16:9, como el PDF de Beamer
        presentacion.slide_height = Emu(6858000)
        vacia = presentacion.slide_layouts[6]
        ancho, alto = presentacion.slide_width, presentacion.slide_height
        for indice, pagina in enumerate(paginas):
            diapositiva = presentacion.slides.add_slide(vacia)
            diapositiva.shapes.add_picture(str(pagina), 0, 0, width=ancho, height=alto)
            if indice in cajas:
                nombre, (x0, y0, x1, y1) = cajas[indice]
                video, _ = VIDEOS[nombre]
                diapositiva.shapes.add_movie(
                    str(video), Emu(int(x0 * ancho)), Emu(int(y0 * alto)),
                    Emu(int((x1 - x0) * ancho)), Emu(int((y1 - y0) * alto)),
                    poster_frame_image=str(PRESENTACION / "figuras" / nombre),
                    mime_type="video/mp4",
                )
                print(f"diapositiva {indice + 1}: {video.name}")
        presentacion.save(salida)
    print(f"{salida.name}: {len(paginas)} diapositivas, {salida.stat().st_size / 1e6:.1f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
