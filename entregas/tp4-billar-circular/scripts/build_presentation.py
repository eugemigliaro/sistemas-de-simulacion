#!/usr/bin/env python3
"""Compila la presentación: PDF para entregar y PPTX con los videos para exponer.

El PDF lleva un fotograma y el enlace de cada animación; en vivo las
animaciones tienen que estar embebidas [GPRES, pp. 2-3]. Cada diapositiva del
PPTX es la página del PDF como imagen y, sobre las que tienen animación, se
superpone el MP4 donde está el fotograma. Para ubicarlo se compila una copia
con cada fotograma reemplazado por un color propio y se busca ese color.
"""
import json
import re
import shutil
import subprocess
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from common import ROOT, VIDEOS

FOLDER = ROOT / "presentacion"
BUILD = FOLDER / ".build"
STEM = "SdS_TP4_2026Q2G07CS_Presentacion"
OFFICIAL = ROOT / "entrega-final" / "SdS_TP4_2026Q2G07CS_Presentación.pdf"
WIDTH, HEIGHT = 1920, 1080
MARKERS = {"densidad": (255, 0, 255), "posicion": (0, 255, 255), "rapidez": (0, 255, 0)}
LINK = re.compile(r"https://(www\.)?(youtube\.com|youtu\.be|vimeo\.com)/[^\s{}%\\]+")


def tectonic() -> str:
    found = shutil.which("tectonic")
    candidates = [found, "/opt/homebrew/bin/tectonic", "/usr/lib/chatgpt/resources/tectonic/tectonic"]
    return next(c for c in candidates if c and Path(c).exists())


def write_links() -> None:
    links = json.loads((FOLDER / "enlaces.json").read_text())
    lines = []
    for name in MARKERS:
        url = links.get(name)
        if url and not LINK.fullmatch(url):
            raise ValueError(f"enlace inválido para {name}: {url}")
        text = rf"\url{{{url}}}" if url else "Enlace de YouTube pendiente"
        lines.append(rf"\expandafter\newcommand\csname enlace{name}\endcsname{{{text}}}")
    (FOLDER / "enlaces.tex").write_text("\n".join(lines) + "\n")


def compile_pdf(folder: Path, output: Path) -> Path:
    output.mkdir(parents=True, exist_ok=True)
    result = subprocess.run([tectonic(), "--only-cached", "--keep-logs", "--outdir", str(output), f"{STEM}.tex"],
                            cwd=folder, capture_output=True, text=True)
    if result.returncode != 0:
        raise SystemExit(result.stderr[-3000:])
    return output / f"{STEM}.pdf"


def render(pdf: Path, folder: Path) -> list[Path]:
    shutil.rmtree(folder, ignore_errors=True)
    folder.mkdir(parents=True)
    subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(WIDTH), "-scale-to-y", str(HEIGHT), str(pdf),
                    str(folder / "p")], check=True)
    return sorted(folder.glob("p-*.png"))


def locate_videos() -> dict[int, tuple[str, tuple[float, float, float, float]]]:
    """Página (desde 0) -> (animación, caja en fracciones de la página)."""
    copy = BUILD / "mascara"
    shutil.rmtree(copy, ignore_errors=True)
    shutil.copytree(FOLDER, copy, ignore=shutil.ignore_patterns(".build", "videos", "*.pptx", "*.pdf"))
    for figure in (FOLDER / "figuras").glob("*.pdf"):
        shutil.copy2(figure, copy / "figuras" / figure.name)
    for name, color in MARKERS.items():
        with Image.open(FOLDER / "figuras" / f"anim-{name}.png") as poster:
            Image.new("RGB", poster.size, color).save(copy / "figuras" / f"anim-{name}.png")
    boxes = {}
    for index, page in enumerate(render(compile_pdf(copy, copy / "salida"), BUILD / "mascara-png")):
        pixels = np.asarray(Image.open(page).convert("RGB")).astype(int)
        for name, color in MARKERS.items():
            hit = np.all(np.abs(pixels - np.array(color)) < 12, axis=2)
            if hit.sum() > 5000:
                rows, columns = np.nonzero(hit)
                boxes[index] = (name, (columns.min() / WIDTH, rows.min() / HEIGHT,
                                       (columns.max() + 1) / WIDTH, (rows.max() + 1) / HEIGHT))
    missing = set(MARKERS) - {name for name, _ in boxes.values()}
    if missing:
        raise SystemExit(f"no se encontró el fotograma de {sorted(missing)}")
    return boxes


def build_pptx(pages: list[Path], boxes: dict) -> Path:
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Emu(12192000), Emu(6858000)
    width, height = presentation.slide_width, presentation.slide_height
    for index, page in enumerate(pages):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(page), 0, 0, width=width, height=height)
        if index in boxes:
            name, (x0, y0, x1, y1) = boxes[index]
            slide.shapes.add_movie(str(VIDEOS / f"{name}.mp4"), Emu(int(x0 * width)), Emu(int(y0 * height)),
                                   Emu(int((x1 - x0) * width)), Emu(int((y1 - y0) * height)),
                                   poster_frame_image=str(FOLDER / "figuras" / f"anim-{name}.png"),
                                   mime_type="video/mp4")
    output = FOLDER / f"{STEM}.pptx"
    presentation.save(output)
    return output


def main() -> None:
    write_links()
    pdf = compile_pdf(FOLDER, BUILD)
    OFFICIAL.parent.mkdir(exist_ok=True)
    shutil.copy2(pdf, OFFICIAL)
    pages = render(pdf, BUILD / "paginas")
    pptx = build_pptx(pages, locate_videos())
    print(f"{OFFICIAL.relative_to(ROOT)}: {len(pages)} diapositivas")
    print(f"{pptx.relative_to(ROOT)}: {pptx.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
